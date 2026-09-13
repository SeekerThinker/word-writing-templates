#!/usr/bin/env python3
"""Generate Windows/macOS Word .dotx writing templates from code only.

The builder creates six numbering schemes for each platform. Each scheme has:
- a minimal starter that preserves the one-click beginner workflow;
- a structured starter with common front/back matter for users who want it.

The builder first asks python-docx to create a standards-compliant Word package,
then patches the few OOXML pieces that python-docx does not expose directly:
multilevel numbering, .dotx content type, and Word keyboard customizations.
"""
from __future__ import annotations

import argparse
import shutil
import tempfile
import zipfile
from pathlib import Path

from lxml import etree
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]

W = 'http://schemas.openxmlformats.org/wordprocessingml/2006/main'
WNE = 'http://schemas.microsoft.com/office/word/2006/wordml'
CT = 'http://schemas.openxmlformats.org/package/2006/content-types'
PR = 'http://schemas.openxmlformats.org/package/2006/relationships'
R = 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
NS = {'w': W, 'wne': WNE, 'ct': CT, 'pr': PR}
Q = lambda ns, tag: f'{{{ns}}}{tag}'

PLATFORMS = {
    'windows': {'label': 'Windows', 'fonts': {'body': '宋体', 'heading': '黑体', 'quote': '楷体', 'title': '微软雅黑'}},
    'macos': {'label': 'macOS', 'fonts': {'body': 'Songti SC', 'heading': 'PingFang SC', 'quote': 'Kaiti SC', 'title': 'PingFang SC'}},
}

TEMPLATES = {
    'book-cn-traditional': {'kind': 'books', 'label': '书籍-中文传统', 'title': '在这里输入书名', 'h1': '在这里输入章标题', 'levels': [('chineseCountingThousand', '第%1章'), ('chineseCountingThousand', '第%2节'), ('chineseCountingThousand', '%3、'), ('chineseCountingThousand', '（%4）')]},
    'book-chapter-decimal': {'kind': 'books', 'label': '书籍-章节数字', 'title': '在这里输入书名', 'h1': '在这里输入章标题', 'levels': [('decimal', '第%1章'), ('decimal', '%1.%2'), ('decimal', '%1.%2.%3'), ('decimal', '%1.%2.%3.%4')]},
    'book-pure-decimal': {'kind': 'books', 'label': '书籍-纯数字', 'title': '在这里输入书名', 'h1': '在这里输入章标题', 'levels': [('decimal', '%1'), ('decimal', '%1.%2'), ('decimal', '%1.%2.%3'), ('decimal', '%1.%2.%3.%4')]},
    'article-cn-academic': {'kind': 'articles', 'label': '文章-中文论文', 'title': '在这里输入文章标题', 'h1': '在这里输入一级标题', 'levels': [('chineseCountingThousand', '%1、'), ('chineseCountingThousand', '（%2）'), ('decimal', '%3.'), ('decimal', '（%4）')]},
    'article-decimal': {'kind': 'articles', 'label': '文章-数字层级', 'title': '在这里输入文章标题', 'h1': '在这里输入一级标题', 'levels': [('decimal', '%1'), ('decimal', '%1.%2'), ('decimal', '%1.%2.%3'), ('decimal', '%1.%2.%3.%4')]},
    'article-cn-compact': {'kind': 'articles', 'label': '文章-中文简洁', 'title': '在这里输入文章标题', 'h1': '在这里输入一级标题', 'levels': [('chineseCountingThousand', '%1、'), ('decimal', '%2.'), ('decimal', '（%3）'), ('decimalEnclosedCircle', '%4')]},
}

KEYMAP_XML = f'''<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<wne:tcg xmlns:r="{R}" xmlns:wne="{WNE}"><wne:keymaps>
<wne:keymap wne:kcmPrimary="0631"><wne:acd wne:acdName="acd0"/></wne:keymap>
<wne:keymap wne:kcmPrimary="0632"><wne:acd wne:acdName="acd1"/></wne:keymap>
<wne:keymap wne:kcmPrimary="0633"><wne:acd wne:acdName="acd4"/></wne:keymap>
<wne:keymap wne:kcmPrimary="0634"><wne:acd wne:acdName="acd5"/></wne:keymap>
<wne:keymap wne:kcmPrimary="0642"><wne:acd wne:acdName="acd6"/></wne:keymap>
<wne:keymap wne:kcmPrimary="0651"><wne:acd wne:acdName="acd2"/></wne:keymap>
<wne:keymap wne:kcmPrimary="0646"><wne:fci wne:fciName="InsertFootnoteNow" wne:swArg="0000"/></wne:keymap>
<wne:keymap wne:kcmPrimary="065A"><wne:acd wne:acdName="acd3"/></wne:keymap>
</wne:keymaps><wne:toolbars><wne:acdManifest>
<wne:acdEntry wne:acdName="acd0"/><wne:acdEntry wne:acdName="acd1"/><wne:acdEntry wne:acdName="acd2"/><wne:acdEntry wne:acdName="acd3"/><wne:acdEntry wne:acdName="acd4"/><wne:acdEntry wne:acdName="acd5"/><wne:acdEntry wne:acdName="acd6"/>
</wne:acdManifest></wne:toolbars><wne:acds>
<wne:acd wne:argValue="AQAAAAEA" wne:acdName="acd0" wne:fciIndexBasedOn="0065"/>
<wne:acd wne:argValue="AQAAAAIA" wne:acdName="acd1" wne:fciIndexBasedOn="0065"/>
<wne:acd wne:argValue="AgBBAEMA" wne:acdName="acd2" wne:fciIndexBasedOn="0065"/>
<wne:acd wne:argValue="AQAAAEIA" wne:acdName="acd3" wne:fciIndexBasedOn="0065"/>
<wne:acd wne:argValue="AQAAAAMA" wne:acdName="acd4" wne:fciIndexBasedOn="0065"/>
<wne:acd wne:argValue="AQAAAAQA" wne:acdName="acd5" wne:fciIndexBasedOn="0065"/>
<wne:acd wne:argValue="AgBoiDxo" wne:acdName="acd6" wne:fciIndexBasedOn="0065"/>
</wne:acds></wne:tcg>'''.encode('utf-8')


def set_font(style, family, size=None, bold=None):
    style.font.name = family
    if size is not None:
        style.font.size = Pt(size)
    if bold is not None:
        style.font.bold = bold
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement('w:rFonts')
        rpr.insert(0, rfonts)
    for attr in ('ascii', 'eastAsia', 'hAnsi', 'cs'):
        rfonts.set(qn(f'w:{attr}'), family)


def mark_quick_style(style):
    if style.element.find(qn('w:qFormat')) is None:
        style.element.append(OxmlElement('w:qFormat'))


def set_outline_level(style, level: int):
    ppr = style.element.get_or_add_pPr()
    old = ppr.find(qn('w:outlineLvl'))
    if old is not None:
        ppr.remove(old)
    node = OxmlElement('w:outlineLvl')
    node.set(qn('w:val'), str(level))
    ppr.append(node)


def get_or_add_numpr(ppr, ilvl, num_id='1'):
    for old in ppr.findall(qn('w:numPr')):
        ppr.remove(old)
    numpr = OxmlElement('w:numPr')
    il = OxmlElement('w:ilvl')
    il.set(qn('w:val'), str(ilvl))
    numpr.append(il)
    ni = OxmlElement('w:numId')
    ni.set(qn('w:val'), str(num_id))
    numpr.append(ni)
    ppr.insert(0, numpr)


def make_numbering_xml(spec, fonts):
    root = etree.Element(Q(W, 'numbering'), nsmap={'w': W})
    abstract = etree.SubElement(root, Q(W, 'abstractNum'))
    abstract.set(Q(W, 'abstractNumId'), '1')
    etree.SubElement(abstract, Q(W, 'nsid')).set(Q(W, 'val'), '4A6C6F6E')
    etree.SubElement(abstract, Q(W, 'multiLevelType')).set(Q(W, 'val'), 'multilevel')
    etree.SubElement(abstract, Q(W, 'tmpl')).set(Q(W, 'val'), '57545432')
    style_ids = ['Heading1', 'Heading2', 'Heading3', 'Heading4']
    families = [fonts['heading'], fonts['heading'], fonts['quote'], fonts['body']]
    sizes = [28, 26, 24, 24]
    for i, ((fmt, text), sid, family, size) in enumerate(zip(spec['levels'], style_ids, families, sizes)):
        lvl = etree.SubElement(abstract, Q(W, 'lvl'))
        lvl.set(Q(W, 'ilvl'), str(i))
        etree.SubElement(lvl, Q(W, 'start')).set(Q(W, 'val'), '1')
        if i > 0:
            etree.SubElement(lvl, Q(W, 'lvlRestart')).set(Q(W, 'val'), str(i))
        etree.SubElement(lvl, Q(W, 'numFmt')).set(Q(W, 'val'), fmt)
        etree.SubElement(lvl, Q(W, 'pStyle')).set(Q(W, 'val'), sid)
        etree.SubElement(lvl, Q(W, 'suff')).set(Q(W, 'val'), 'space')
        etree.SubElement(lvl, Q(W, 'lvlText')).set(Q(W, 'val'), text)
        etree.SubElement(lvl, Q(W, 'lvlJc')).set(Q(W, 'val'), 'left')
        ppr = etree.SubElement(lvl, Q(W, 'pPr'))
        ind = etree.SubElement(ppr, Q(W, 'ind'))
        ind.set(Q(W, 'left'), str(i * 284))
        ind.set(Q(W, 'firstLine'), '0')
        rpr = etree.SubElement(lvl, Q(W, 'rPr'))
        rf = etree.SubElement(rpr, Q(W, 'rFonts'))
        for attr in ('ascii', 'eastAsia', 'hAnsi', 'cs'):
            rf.set(Q(W, attr), family)
        etree.SubElement(rpr, Q(W, 'sz')).set(Q(W, 'val'), str(size))
        etree.SubElement(rpr, Q(W, 'szCs')).set(Q(W, 'val'), str(size))
    num = etree.SubElement(root, Q(W, 'num'))
    num.set(Q(W, 'numId'), '1')
    etree.SubElement(num, Q(W, 'abstractNumId')).set(Q(W, 'val'), '1')
    return etree.tostring(root, xml_declaration=True, encoding='UTF-8', standalone=True)


def ensure_style(doc, name):
    try:
        return doc.styles[name]
    except KeyError:
        return doc.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)


def style_doc(doc, platform):
    f = PLATFORMS[platform]['fonts']
    sec = doc.sections[0]
    sec.page_width = Mm(210)
    sec.page_height = Mm(297)
    sec.top_margin = Mm(25.4)
    sec.bottom_margin = Mm(25.4)
    sec.left_margin = Mm(25.4)
    sec.right_margin = Mm(25.4)

    normal = doc.styles['Normal']
    set_font(normal, f['body'], 12)
    normal.paragraph_format.line_spacing = 1.5

    body = doc.styles['Body Text']
    set_font(body, f['body'], 12)
    body.paragraph_format.line_spacing = 1.5
    body.paragraph_format.first_line_indent = Mm(8)
    body.paragraph_format.space_after = Pt(4)

    title = doc.styles['Title']
    set_font(title, f['title'], 20, True)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(8)

    for i, (size, before, after, family) in enumerate(
        [(14, 12, 8, f['heading']), (13, 8, 5, f['heading']), (12, 6, 4, f['quote']), (12, 4, 2, f['body'])],
        start=1,
    ):
        st = doc.styles[f'Heading {i}']
        set_font(st, family, size, i < 3)
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.line_spacing = 1.5
        get_or_add_numpr(st.element.get_or_add_pPr(), i - 1, '1')

    quote = ensure_style(doc, 'AC')
    set_font(quote, f['quote'], 12)
    quote.paragraph_format.left_indent = Mm(7)
    quote.paragraph_format.right_indent = Mm(7)
    quote.paragraph_format.line_spacing = 1.5

    table = ensure_style(doc, '表格')
    set_font(table, f['body'], 10)
    table.paragraph_format.first_line_indent = Mm(0)
    table.paragraph_format.space_after = Pt(0)
    table.paragraph_format.line_spacing = 1

    foot = ensure_style(doc, 'Footnote Text')
    set_font(foot, f['body'], 9)
    foot.paragraph_format.line_spacing = 1

    structure = ensure_style(doc, '结构标题')
    structure.base_style = normal
    set_font(structure, f['heading'], 14, True)
    structure.paragraph_format.space_before = Pt(12)
    structure.paragraph_format.space_after = Pt(6)
    structure.paragraph_format.keep_with_next = True
    structure.paragraph_format.first_line_indent = Mm(0)
    set_outline_level(structure, 0)
    mark_quick_style(structure)

    toc_title = ensure_style(doc, '目录标题')
    toc_title.base_style = normal
    set_font(toc_title, f['heading'], 14, True)
    toc_title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    toc_title.paragraph_format.space_before = Pt(12)
    toc_title.paragraph_format.space_after = Pt(8)
    toc_title.paragraph_format.first_line_indent = Mm(0)

    abstract = ensure_style(doc, '摘要')
    abstract.base_style = normal
    set_font(abstract, f['body'], 11)
    abstract.paragraph_format.first_line_indent = Mm(0)
    abstract.paragraph_format.line_spacing = 1.5
    abstract.paragraph_format.space_after = Pt(5)
    mark_quick_style(abstract)

    keywords = ensure_style(doc, '关键词')
    keywords.base_style = normal
    set_font(keywords, f['body'], 11)
    keywords.paragraph_format.first_line_indent = Mm(0)
    keywords.paragraph_format.space_after = Pt(8)

    refs = ensure_style(doc, '参考文献')
    refs.base_style = normal
    set_font(refs, f['body'], 10.5)
    refs.paragraph_format.left_indent = Mm(7)
    refs.paragraph_format.first_line_indent = Mm(-7)
    refs.paragraph_format.line_spacing = 1.25
    refs.paragraph_format.space_after = Pt(3)
    mark_quick_style(refs)

    note = ensure_style(doc, '模板提示')
    note.base_style = normal
    set_font(note, f['body'], 9)
    note.font.color.rgb = RGBColor(110, 110, 110)
    note.paragraph_format.first_line_indent = Mm(0)
    note.paragraph_format.space_after = Pt(4)


def add_toc(paragraph):
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'TOC \\o "1-3" \\h \\z \\u')
    run = OxmlElement('w:r')
    text = OxmlElement('w:t')
    text.text = '目录会根据标题自动生成；如未刷新，请在目录上右键选择“更新域”。'
    run.append(text)
    fld.append(run)
    paragraph._p.append(fld)


def enable_field_updates(doc):
    settings = doc.settings.element
    current = settings.find(qn('w:updateFields'))
    if current is None:
        current = OxmlElement('w:updateFields')
        settings.append(current)
    current.set(qn('w:val'), 'true')


def add_minimal_content(doc, spec):
    doc.add_paragraph(spec['title'], style='Title')
    doc.add_paragraph(spec['h1'], style='Heading 1')
    doc.add_paragraph('从这里开始写作。', style='Body Text')


def add_structured_book(doc, spec):
    doc.add_paragraph(spec['title'], style='Title')
    doc.add_paragraph('前言（可选）', style='结构标题')
    doc.add_paragraph('在这里写前言；不需要前言时，可以直接删除这一节。', style='Body Text')
    doc.add_paragraph('目录', style='目录标题')
    toc = doc.add_paragraph(style='模板提示')
    add_toc(toc)
    doc.add_page_break()
    doc.add_paragraph(spec['h1'], style='Heading 1')
    doc.add_paragraph('从这里开始写正文。需要新章节时，继续使用“标题 1”即可。', style='Body Text')
    doc.add_page_break()
    doc.add_paragraph('附录（可选）', style='结构标题')
    doc.add_paragraph('需要附录时从这里开始；不需要可以删除这一节。', style='Body Text')
    doc.add_paragraph('参考文献（可选）', style='结构标题')
    doc.add_paragraph('[1] 在这里填写参考文献。', style='参考文献')


def add_structured_article(doc, spec):
    doc.add_paragraph(spec['title'], style='Title')
    doc.add_paragraph('摘要（可选）', style='结构标题')
    doc.add_paragraph('在这里填写摘要；一般文章不需要时可以删除这一节。', style='摘要')
    doc.add_paragraph('关键词：在这里填写关键词。', style='关键词')
    doc.add_paragraph(spec['h1'], style='Heading 1')
    doc.add_paragraph('从这里开始写正文。', style='Body Text')
    doc.add_paragraph('参考文献（可选）', style='结构标题')
    doc.add_paragraph('[1] 在这里填写参考文献。', style='参考文献')


def make_docx(spec, platform, preview, path, structured=False):
    doc = Document()
    style_doc(doc, platform)
    enable_field_updates(doc)
    if preview:
        doc.add_paragraph(spec['label'], style='Title')
        doc.add_paragraph('一级标题示例', style='Heading 1')
        doc.add_paragraph('这是一段正文示例。标题编号会随着章节层级自动变化。', style='Body Text')
        doc.add_paragraph('二级标题示例', style='Heading 2')
        doc.add_paragraph('正文示例。', style='Body Text')
        doc.add_paragraph('三级标题示例', style='Heading 3')
        doc.add_paragraph('正文示例。', style='Body Text')
        doc.add_paragraph('四级标题示例', style='Heading 4')
        doc.add_paragraph('正文示例。', style='Body Text')
    elif structured and spec['kind'] == 'books':
        add_structured_book(doc, spec)
    elif structured and spec['kind'] == 'articles':
        add_structured_article(doc, spec)
    else:
        add_minimal_content(doc, spec)
    variant = '常用结构' if structured else '直接开始'
    doc.core_properties.title = f"{spec['label']} - {variant} - {PLATFORMS[platform]['label']}"
    doc.core_properties.subject = 'Word 结构化写作模板'
    doc.core_properties.author = ''
    doc.core_properties.last_modified_by = ''
    doc.save(path)


def patch_to_dotx(docx_path, out, spec, platform):
    with zipfile.ZipFile(docx_path, 'r') as zin:
        files = {name: zin.read(name) for name in zin.namelist()}
    files['word/numbering.xml'] = make_numbering_xml(spec, PLATFORMS[platform]['fonts'])
    ct = etree.fromstring(files['[Content_Types].xml'])
    for ov in ct.findall('ct:Override', NS):
        if ov.get('PartName') == '/word/document.xml':
            ov.set('ContentType', 'application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml')
    if not any(ov.get('PartName') == '/word/customizations.xml' for ov in ct.findall('ct:Override', NS)):
        ov = etree.SubElement(ct, Q(CT, 'Override'))
        ov.set('PartName', '/word/customizations.xml')
        ov.set('ContentType', 'application/vnd.ms-word.keyMapCustomizations+xml')
    files['[Content_Types].xml'] = etree.tostring(ct, xml_declaration=True, encoding='UTF-8', standalone=True)
    rels = etree.fromstring(files['word/_rels/document.xml.rels'])
    if not any(rel.get('Type', '').endswith('/keyMapCustomizations') for rel in rels):
        used = {rel.get('Id') for rel in rels}
        n = 1
        while f'rId{n}' in used:
            n += 1
        rel = etree.SubElement(rels, Q(PR, 'Relationship'))
        rel.set('Id', f'rId{n}')
        rel.set('Type', 'http://schemas.microsoft.com/office/2006/relationships/keyMapCustomizations')
        rel.set('Target', 'customizations.xml')
    files['word/_rels/document.xml.rels'] = etree.tostring(rels, xml_declaration=True, encoding='UTF-8', standalone=True)
    files['word/customizations.xml'] = KEYMAP_XML
    out.parent.mkdir(parents=True, exist_ok=True)
    with zipfile.ZipFile(out, 'w', zipfile.ZIP_DEFLATED) as zout:
        for name, data in files.items():
            zout.writestr(name, data)


def build_one(out, spec, platform, preview=False, structured=False):
    with tempfile.TemporaryDirectory() as td:
        tmp = Path(td) / 'base.docx'
        make_docx(spec, platform, preview, tmp, structured=structured)
        patch_to_dotx(tmp, out, spec, platform)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument('--preview', action='store_true')
    args = ap.parse_args()
    if (ROOT / 'templates').exists():
        shutil.rmtree(ROOT / 'templates')
    if args.preview and (ROOT / 'build' / 'previews').exists():
        shutil.rmtree(ROOT / 'build' / 'previews')
    for platform in PLATFORMS:
        for key, spec in TEMPLATES.items():
            simple = ROOT / 'templates' / platform / spec['kind'] / f"{spec['label']}.dotx"
            build_one(simple, spec, platform)
            print(simple.relative_to(ROOT))
            structured = ROOT / 'templates' / platform / 'structured' / spec['kind'] / f"{spec['label']}-常用结构.dotx"
            build_one(structured, spec, platform, structured=True)
            print(structured.relative_to(ROOT))
            if args.preview:
                build_one(ROOT / 'build' / 'previews' / platform / f'{key}.dotx', spec, platform, True)


if __name__ == '__main__':
    main()
