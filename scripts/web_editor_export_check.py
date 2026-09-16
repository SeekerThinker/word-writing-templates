#!/usr/bin/env python3
"""Validate the web editor's Word-export assumptions against every real template.

The check mirrors the browser export package boundary: real Heading 1–4 styles,
numbering, a genuine Word table using the template's table-text paragraph style,
and a genuine Word footnote part + relationship + reference. The generated .docx
must remain readable by python-docx and preserve template shortcut customizations.
"""
from __future__ import annotations

import hashlib
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from lxml import etree

ROOT = Path(__file__).resolve().parents[1]
TEMPLATES = ROOT / "templates"

W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
CT = "http://schemas.openxmlformats.org/package/2006/content-types"
PR = "http://schemas.openxmlformats.org/package/2006/relationships"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS = {"w": W, "ct": CT, "pr": PR}
Q = lambda ns, tag: f"{{{ns}}}{tag}"

TEMPLATE_MAIN = "application/vnd.openxmlformats-officedocument.wordprocessingml.template.main+xml"
DOCUMENT_MAIN = "application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"
FOOTNOTE_CT = "application/vnd.openxmlformats-officedocument.wordprocessingml.footnotes+xml"
FOOTNOTE_REL = f"{R}/footnotes"

EXPECTED = {
    "books": ["书籍-中文传统.dotx", "书籍-章节数字.dotx", "书籍-纯数字.dotx"],
    "articles": ["文章-中文论文.dotx", "文章-数字层级.dotx", "文章-中文简洁.dotx"],
}

REQUIRED_PARTS = {
    "[Content_Types].xml",
    "word/document.xml",
    "word/styles.xml",
    "word/numbering.xml",
    "word/settings.xml",
    "word/customizations.xml",
    "word/_rels/document.xml.rels",
}


def parse(data: bytes) -> etree._Element:
    return etree.fromstring(data)


def text_paragraph(style_id: str, text: str) -> etree._Element:
    p = etree.Element(Q(W, "p"), nsmap={"w": W})
    ppr = etree.SubElement(p, Q(W, "pPr"))
    pstyle = etree.SubElement(ppr, Q(W, "pStyle"))
    pstyle.set(Q(W, "val"), style_id)
    run = etree.SubElement(p, Q(W, "r"))
    node = etree.SubElement(run, Q(W, "t"))
    node.text = text
    return p


def footnote_reference_paragraph(style_id: str, before: str, footnote_id: int, after: str) -> etree._Element:
    p = etree.Element(Q(W, "p"), nsmap={"w": W})
    ppr = etree.SubElement(p, Q(W, "pPr"))
    pstyle = etree.SubElement(ppr, Q(W, "pStyle"))
    pstyle.set(Q(W, "val"), style_id)
    run = etree.SubElement(p, Q(W, "r"))
    etree.SubElement(run, Q(W, "t")).text = before
    ref_run = etree.SubElement(p, Q(W, "r"))
    rpr = etree.SubElement(ref_run, Q(W, "rPr"))
    valign = etree.SubElement(rpr, Q(W, "vertAlign"))
    valign.set(Q(W, "val"), "superscript")
    ref = etree.SubElement(ref_run, Q(W, "footnoteReference"))
    ref.set(Q(W, "id"), str(footnote_id))
    run2 = etree.SubElement(p, Q(W, "r"))
    etree.SubElement(run2, Q(W, "t")).text = after
    return p


def find_style_id(styles_data: bytes, style_name: str, fallback: str) -> str:
    root = parse(styles_data)
    for style in root.findall("w:style", NS):
        name = style.find("w:name", NS)
        if name is not None and name.get(Q(W, "val")) == style_name:
            return style.get(Q(W, "styleId")) or fallback
    return fallback


def sample_table(style_id: str) -> etree._Element:
    tbl = etree.Element(Q(W, "tbl"), nsmap={"w": W})
    tblpr = etree.SubElement(tbl, Q(W, "tblPr"))
    width = etree.SubElement(tblpr, Q(W, "tblW"))
    width.set(Q(W, "w"), "0")
    width.set(Q(W, "type"), "auto")
    borders = etree.SubElement(tblpr, Q(W, "tblBorders"))
    for side in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = etree.SubElement(borders, Q(W, side))
        border.set(Q(W, "val"), "single")
        border.set(Q(W, "sz"), "4")
        border.set(Q(W, "color"), "BFBFBF")
    grid = etree.SubElement(tbl, Q(W, "tblGrid"))
    for _ in range(2):
        col = etree.SubElement(grid, Q(W, "gridCol"))
        col.set(Q(W, "w"), "4500")
    for values in (("A1", "B1"), ("A2", "B2")):
        tr = etree.SubElement(tbl, Q(W, "tr"))
        for value in values:
            tc = etree.SubElement(tr, Q(W, "tc"))
            tc.append(text_paragraph(style_id, value))
    return tbl


def patch_document(data: bytes, styles_data: bytes) -> bytes:
    root = parse(data)
    body = root.find("w:body", NS)
    assert body is not None, "word/document.xml has no body"
    sectpr = body.find("w:sectPr", NS)
    preserved = etree.fromstring(etree.tostring(sectpr)) if sectpr is not None else None
    for child in list(body):
        body.remove(child)
    body.append(text_paragraph("Title", "Web editor export check"))
    for level in range(1, 5):
        body.append(text_paragraph(f"Heading{level}", f"Heading level {level}"))
        body.append(text_paragraph("BodyText", f"Body text under heading {level}."))
    table_style = find_style_id(styles_data, "表格", "BodyText")
    body.append(sample_table(table_style))
    body.append(footnote_reference_paragraph("BodyText", "Text with footnote", 1, "."))
    if preserved is not None:
        body.append(preserved)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def make_footnotes(styles_data: bytes) -> bytes:
    foot_style = find_style_id(styles_data, "Footnote Text", "FootnoteText")
    root = etree.Element(Q(W, "footnotes"), nsmap={"w": W})
    sep = etree.SubElement(root, Q(W, "footnote"))
    sep.set(Q(W, "type"), "separator")
    sep.set(Q(W, "id"), "-1")
    p = etree.SubElement(sep, Q(W, "p"))
    r = etree.SubElement(p, Q(W, "r"))
    etree.SubElement(r, Q(W, "separator"))
    cont = etree.SubElement(root, Q(W, "footnote"))
    cont.set(Q(W, "type"), "continuationSeparator")
    cont.set(Q(W, "id"), "0")
    p = etree.SubElement(cont, Q(W, "p"))
    r = etree.SubElement(p, Q(W, "r"))
    etree.SubElement(r, Q(W, "continuationSeparator"))
    fn = etree.SubElement(root, Q(W, "footnote"))
    fn.set(Q(W, "id"), "1")
    p = text_paragraph(foot_style, "")
    first_run = p.find("w:r", NS)
    assert first_run is not None
    p.remove(first_run)
    ref_run = etree.SubElement(p, Q(W, "r"))
    etree.SubElement(ref_run, Q(W, "footnoteRef"))
    text_run = etree.SubElement(p, Q(W, "r"))
    t = etree.SubElement(text_run, Q(W, "t"))
    t.text = " Real footnote from browser model."
    fn.append(p)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def patch_content_types(data: bytes) -> bytes:
    root = parse(data)
    main = None
    for override in root.findall("ct:Override", NS):
        if override.get("PartName") == "/word/document.xml":
            main = override
            break
    assert main is not None, "missing /word/document.xml content type override"
    assert main.get("ContentType") == TEMPLATE_MAIN, "source package is not a .dotx template"
    main.set("ContentType", DOCUMENT_MAIN)
    footnote_override = next((o for o in root.findall("ct:Override", NS) if o.get("PartName") == "/word/footnotes.xml"), None)
    if footnote_override is None:
        footnote_override = etree.SubElement(root, Q(CT, "Override"))
        footnote_override.set("PartName", "/word/footnotes.xml")
    footnote_override.set("ContentType", FOOTNOTE_CT)
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def patch_document_rels(data: bytes) -> bytes:
    root = parse(data)
    existing = root.xpath("pr:Relationship[@Type=$t]", namespaces=NS, t=FOOTNOTE_REL)
    if not existing:
        max_id = 0
        for rel in root.findall("pr:Relationship", NS):
            value = rel.get("Id", "")
            if value.startswith("rId") and value[3:].isdigit():
                max_id = max(max_id, int(value[3:]))
        rel = etree.SubElement(root, Q(PR, "Relationship"))
        rel.set("Id", f"rId{max_id + 1}")
        rel.set("Type", FOOTNOTE_REL)
        rel.set("Target", "footnotes.xml")
    return etree.tostring(root, xml_declaration=True, encoding="UTF-8", standalone=True)


def assert_styles(data: bytes) -> None:
    root = parse(data)
    for level in range(1, 5):
        style_id = f"Heading{level}"
        style = root.xpath(f"w:style[@w:styleId='{style_id}']", namespaces=NS)
        assert style, f"missing {style_id} style"
        num_id = style[0].xpath("w:pPr/w:numPr/w:numId/@w:val", namespaces=NS)
        assert num_id == ["1"], f"{style_id} must use numbering numId=1, got {num_id}"
    assert find_style_id(data, "表格", "") != "", "missing table-text style 表格"
    assert find_style_id(data, "Footnote Text", "") != "", "missing Footnote Text style"


def assert_numbering(data: bytes) -> None:
    root = parse(data)
    nums = root.xpath("w:num[@w:numId='1']/w:abstractNumId/@w:val", namespaces=NS)
    assert nums, "missing heading numbering numId=1"
    abstract_id = nums[0]
    abstract = root.xpath(f"w:abstractNum[@w:abstractNumId='{abstract_id}']", namespaces=NS)
    assert abstract, f"missing abstract numbering {abstract_id}"
    styles = abstract[0].xpath("w:lvl/w:pStyle/@w:val", namespaces=NS)
    assert styles[:4] == ["Heading1", "Heading2", "Heading3", "Heading4"], f"heading numbering must bind Heading1–Heading4, got {styles[:4]}"


def convert_sample(source: Path, target: Path) -> str:
    with zipfile.ZipFile(source) as zin:
        styles_data = zin.read("word/styles.xml")
        customizations = zin.read("word/customizations.xml")
        custom_hash = hashlib.sha256(customizations).hexdigest()
        with zipfile.ZipFile(target, "w", compression=zipfile.ZIP_DEFLATED) as zout:
            for info in zin.infolist():
                if info.filename == "word/footnotes.xml":
                    continue
                data = zin.read(info.filename)
                if info.filename == "word/document.xml":
                    data = patch_document(data, styles_data)
                elif info.filename == "[Content_Types].xml":
                    data = patch_content_types(data)
                elif info.filename == "word/_rels/document.xml.rels":
                    data = patch_document_rels(data)
                zout.writestr(info, data)
            zout.writestr("word/footnotes.xml", make_footnotes(styles_data))
    return custom_hash


def validate_template(path: Path) -> None:
    with zipfile.ZipFile(path) as zf:
        names = set(zf.namelist())
        missing = REQUIRED_PARTS - names
        assert not missing, f"{path}: missing package parts: {sorted(missing)}"
        assert_styles(zf.read("word/styles.xml"))
        assert_numbering(zf.read("word/numbering.xml"))
        doc = parse(zf.read("word/document.xml"))
        assert doc.find(".//w:body/w:sectPr", NS) is not None, f"{path}: missing section properties"

    with tempfile.TemporaryDirectory() as tmp:
        output = Path(tmp) / "sample.docx"
        customization_hash = convert_sample(path, output)
        document = Document(output)
        paragraphs = [p for p in document.paragraphs if p.text]
        expected_prefix = [
            "Title",
            "Heading 1", "Body Text",
            "Heading 2", "Body Text",
            "Heading 3", "Body Text",
            "Heading 4", "Body Text",
        ]
        actual_prefix = [p.style.name for p in paragraphs[: len(expected_prefix)]]
        assert actual_prefix == expected_prefix, f"{path}: exported styles differ: {actual_prefix}"
        assert len(document.tables) == 1, f"{path}: exported table is not a real Word table"
        assert document.tables[0].cell(0, 0).text == "A1"
        assert document.tables[0].cell(1, 1).text == "B2"
        with zipfile.ZipFile(output) as zf:
            content_types = parse(zf.read("[Content_Types].xml"))
            main = content_types.xpath("ct:Override[@PartName='/word/document.xml']/@ContentType", namespaces=NS)
            assert main == [DOCUMENT_MAIN], f"{path}: export did not become a .docx main part"
            foot_ct = content_types.xpath("ct:Override[@PartName='/word/footnotes.xml']/@ContentType", namespaces=NS)
            assert foot_ct == [FOOTNOTE_CT], f"{path}: missing footnote content type"
            rels = parse(zf.read("word/_rels/document.xml.rels"))
            foot_rels = rels.xpath("pr:Relationship[@Type=$t]/@Target", namespaces=NS, t=FOOTNOTE_REL)
            assert foot_rels == ["footnotes.xml"], f"{path}: missing footnote relationship"
            footnotes = parse(zf.read("word/footnotes.xml"))
            note_text = "".join(footnotes.xpath("w:footnote[@w:id='1']//w:t/text()", namespaces=NS))
            assert "Real footnote" in note_text, f"{path}: footnote part has no real note text"
            doc_xml = parse(zf.read("word/document.xml"))
            refs = doc_xml.xpath(".//w:footnoteReference/@w:id", namespaces=NS)
            assert refs == ["1"], f"{path}: document has no real footnote reference"
            after_hash = hashlib.sha256(zf.read("word/customizations.xml")).hexdigest()
            assert after_hash == customization_hash, f"{path}: export unexpectedly changed template shortcuts"


def main() -> None:
    paths: list[Path] = []
    for platform in ("windows", "macos"):
        for kind, filenames in EXPECTED.items():
            for filename in filenames:
                paths.append(TEMPLATES / platform / kind / filename)
    missing = [str(path.relative_to(ROOT)) for path in paths if not path.exists()]
    if missing:
        raise SystemExit(f"missing templates: {missing}")
    for path in paths:
        validate_template(path)
        print(f"OK {path.relative_to(ROOT)}")
    print(f"Web editor Word-export package check OK ({len(paths)} templates, real tables + footnotes)")


if __name__ == "__main__":
    main()
